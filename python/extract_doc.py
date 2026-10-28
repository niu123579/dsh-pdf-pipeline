#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""提取 PDF / PPTX / DOCX 的文本与公式（Office OMML 公式尽量转成 LaTeX）。"""
import json
import os
import re
import sys

# 强制 UTF-8 输出，避免 Windows 控制台默认 GBK 编码导致中文乱码
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(encoding='utf-8')

from docx.oxml.ns import qn


def _linear_text(el):
    """拼接元素下所有 m:t / a:t / w:t 文本（线性回退形式）。"""
    parts = []
    for t in el.iter():
        if t.tag.endswith('}t'):
            parts.append(t.text or '')
    return ''.join(parts)


FUNC_MAP = {
    'sin': '\\sin', 'cos': '\\cos', 'tan': '\\tan', 'cot': '\\cot',
    'sec': '\\sec', 'csc': '\\csc', 'log': '\\log', 'ln': '\\ln',
    'lim': '\\lim', 'max': '\\max', 'min': '\\min', 'det': '\\det',
    'gcd': '\\gcd', 'deg': '\\deg', 'arg': '\\arg', 'exp': '\\exp',
}

NARY_MAP = {
    '\u2211': '\\sum', '\u222b': '\\int', '\u220f': '\\prod',
    '\u2210': '\\coprod', '\u22c3': '\\bigcup', '\u22c2': '\\bigcap',
    '\u22c1': '\\bigvee', '\u22c0': '\\bigwedge', '\u2a01': '\\bigoplus',
    '\u2a02': '\\bigotimes', '\u222e': '\\oint',
}

ACCENT_MAP = {
    '\u0302': '\\hat', '^': '\\hat', '\u0303': '\\tilde', '~': '\\tilde',
    '\u0304': '\\bar', '\u00af': '\\bar', '\u0308': '\\ddot', '\u00a8': '\\ddot',
    '\u0307': '\\dot', '\u20d7': '\\vec', '\u2192': '\\vec',
}


def _children(el):
    if el is None:
        return ''
    return ''.join(_conv(c) for c in el)


def _conv(el):
    tag = el.tag.split('}')[-1]
    if tag == 'r':
        return _linear_text(el)
    if tag == 'f':
        return '\\frac{%s}{%s}' % (_children(el.find(qn('m:num'))), _children(el.find(qn('m:den'))))
    if tag == 'sSup':
        return '%s^{%s}' % (_children(el.find(qn('m:e'))), _children(el.find(qn('m:sup'))))
    if tag == 'sSub':
        return '%s_{%s}' % (_children(el.find(qn('m:e'))), _children(el.find(qn('m:sub'))))
    if tag == 'sSubSup':
        return '%s_{%s}^{%s}' % (
            _children(el.find(qn('m:e'))), _children(el.find(qn('m:sub'))), _children(el.find(qn('m:sup'))))
    if tag == 'rad':
        deg = el.find(qn('m:deg'))
        e = _children(el.find(qn('m:e')))
        if deg is not None and _linear_text(deg):
            return '\\sqrt[%s]{%s}' % (_children(deg), e)
        return '\\sqrt{%s}' % e
    if tag == 'nary':
        return _nary(el)
    if tag == 'acc':
        return _accent(el)
    if tag == 'func':
        return _func(el)
    if tag == 'd':
        return '\\left( %s \\right)' % _children(el)
    if tag == 'm':
        return _matrix(el)
    if tag in ('bar', 'e', 'num', 'den', 'sub', 'sup', 'deg', 'fName'):
        return _children(el)
    return _linear_text(el)


def _nary(el):
    ch = None
    pr = el.find(qn('m:naryPr'))
    if pr is not None:
        chr_el = pr.find(qn('m:chr'))
        if chr_el is not None:
            ch = chr_el.get(qn('m:val'))
    sym = NARY_MAP.get(ch) if ch else None
    if not sym:
        sym = _linear_text(el) or '\\sum'
    lim = ''
    sup = el.find(qn('m:sup'))
    sub = el.find(qn('m:sub'))
    if sup is not None and _linear_text(sup):
        lim += '^{%s}' % _children(sup)
    if sub is not None and _linear_text(sub):
        lim += '_{%s}' % _children(sub)
    return '%s%s %s' % (sym, lim, _children(el.find(qn('m:e'))))


def _accent(el):
    ch = None
    pr = el.find(qn('m:accPr'))
    if pr is not None:
        chr_el = pr.find(qn('m:chr'))
        if chr_el is not None:
            ch = chr_el.get(qn('m:val'))
    acc = ACCENT_MAP.get(ch, '\\hat')
    return '%s{%s}' % (acc, _children(el.find(qn('m:e'))))


def _func(el):
    fname_el = el.find(qn('m:fName'))
    name = _linear_text(fname_el) if fname_el is not None else ''
    latex = FUNC_MAP.get(name) if name else ''
    if not latex:
        latex = '\\operatorname{%s}' % name if name else ''
    return '%s %s' % (latex, _children(el.find(qn('m:e'))))


def _matrix(el):
    rows = []
    for mr in el.findall(qn('m:mr')):
        cells = [_children(e) for e in mr.findall(qn('m:e'))]
        rows.append(' & '.join(cells))
    return '\\begin{matrix} %s \\end{matrix}' % ' \\\\ '.join(rows)


def omml_to_latex(om):
    """把一个 m:oMath 元素转成 LaTeX 字符串。"""
    s = ''.join(_conv(c) for c in om).strip()
    return s if s else _linear_text(om)


def iter_omath(container):
    for el in container.iter():
        if el.tag.endswith('}oMath'):
            yield el


def extract_pptx(fp, include_notes=True):
    from pptx import Presentation
    prs = Presentation(fp)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        formulas = []
        seen = set()
        for shape in slide.shapes:
            if getattr(shape, 'has_text_frame', False):
                _collect_text_frame(shape.text_frame, texts, formulas, seen)
            if getattr(shape, 'has_table', False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        _collect_text_frame(cell.text_frame, texts, formulas, seen)
            if getattr(shape, 'shape_type', None) == 6:  # group
                _walk_group(shape, texts, formulas, seen)
        slides.append({'index': i, 'text': '\n'.join(texts), 'formulas': formulas})

    notes = []
    if include_notes:
        for i, slide in enumerate(prs.slides, 1):
            try:
                if slide.has_notes_slide:
                    nt = slide.notes_slide.notes_text_frame.text
                    if nt.strip():
                        notes.append({'index': i, 'text': nt})
            except Exception:
                pass
    return {'ok': True, 'type': 'pptx', 'slides': slides, 'notes': notes}


def _collect_text_frame(tf, texts, formulas, seen):
    for para in tf.paragraphs:
        t = ''.join(r.text for r in para.runs)
        if t.strip():
            texts.append(t)
        for om in iter_omath(para._p):
            if id(om) not in seen:
                seen.add(id(om))
                formulas.append(omml_to_latex(om))


def _walk_group(shape, texts, formulas, seen):
    try:
        for sub in shape.shapes:
            if getattr(sub, 'has_text_frame', False):
                _collect_text_frame(sub.text_frame, texts, formulas, seen)
            if getattr(sub, 'shape_type', None) == 6:
                _walk_group(sub, texts, formulas, seen)
    except Exception:
        pass


def extract_docx(fp):
    import docx
    d = docx.Document(fp)
    paras = []
    formulas = []
    for para in d.paragraphs:
        if para.text.strip():
            paras.append(para.text)
        for om in iter_omath(para._p):
            formulas.append(omml_to_latex(om))
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip():
                        paras.append(para.text)
    return {'ok': True, 'type': 'docx', 'paragraphs': paras, 'formulas': formulas}


_MATH_RE = re.compile(r'[\u222b\u2211\u220f\u221a\u221e\u2264\u2265\u2260\u2248\u00b1\u00d7\u00f7'
                      r'\u2202\u2207\u2206\u03b1\u03b2\u03b3\u03b4\u03b5\u03b6\u03b7\u03b8\u03bb\u03bc\u03bd'
                      r'\u03be\u03c0\u03c1\u03c3\u03c4\u03c6\u03c7\u03c8\u03c9\u03a9\u2192\u21d2\u21d4\u2208'
                      r'\u2282\u222a\u2229]|\\[a-zA-Z]+')


def extract_pdf(fp):
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz  # 兼容旧版 PyMuPDF
    except ImportError:
        return {
            'ok': False,
            'error': 'PyMuPDF 未安装',
            'hint': '首次提取 PDF 时插件会自动联网安装 pymupdf，若失败请检查网络。',
        }
    doc = fitz.open(fp)
    pages = []
    for i in range(len(doc)):
        page = doc[i]
        text = page.get_text('text')
        lines = [ln for ln in text.splitlines() if ln.strip()]
        suspected = [ln for ln in lines if _MATH_RE.search(ln)]
        pages.append({'index': i + 1, 'text': text, 'suspected_formula_lines': suspected})
    return {
        'ok': True,
        'type': 'pdf',
        'page_count': len(doc),
        'pages': pages,
        'note': 'PDF 中以图片形式渲染的公式无法提取文本，已用启发式标注疑似公式行。',
    }


def main():
    if len(sys.argv) < 2:
        print(json.dumps({'ok': False, 'error': '缺少文件路径参数'}, ensure_ascii=False))
        return
    fp = sys.argv[1]
    include_notes = not (len(sys.argv) > 2 and sys.argv[2] == 'no-notes')
    if not os.path.exists(fp):
        print(json.dumps({'ok': False, 'error': '文件不存在: %s' % fp}, ensure_ascii=False))
        return
    ext = os.path.splitext(fp)[1].lower()
    try:
        if ext == '.pptx':
            result = extract_pptx(fp, include_notes)
        elif ext == '.docx':
            result = extract_docx(fp)
        elif ext == '.pdf':
            result = extract_pdf(fp)
        else:
            result = {'ok': False, 'error': '不支持的文档类型: %s（支持 .pptx / .docx / .pdf）' % ext}
    except Exception as e:  # noqa: BLE001
        import traceback
        result = {'ok': False, 'error': '%s: %s' % (type(e).__name__, e),
                  'detail': traceback.format_exc()[-3000:]}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
